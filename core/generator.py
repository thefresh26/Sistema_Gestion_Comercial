"""
Genera el Acta de Arrendamiento en Word a partir de:
  - los datos del caso (extraídos automáticamente + correcciones manuales
    aplicadas por encima),
  - la plantilla en blanco (word_templates/ACTA_ARRIENDOS.docx),
  - las fotos del inmueble guardadas como documentos en la base de datos.

Devuelve los bytes del .docx final, listos para guardar en la tabla
`documentos` (tipo='acta_generada') y para servir en la descarga.
"""
from __future__ import annotations

import io
import os

from docx import Document
from docx.shared import Mm

RUTA_PLANTILLA = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "word_templates", "ACTA_ARRIENDOS.docx",
)

SIN_IMAGEN = "No hay imagen para este documento"


def _reemplazar_texto_parrafo(parrafo, reemplazos: dict[str, str]) -> None:
    texto = parrafo.text
    cambiado = False
    for marcador, valor in reemplazos.items():
        if marcador in texto:
            texto = texto.replace(marcador, valor or "—")
            cambiado = True
    if cambiado:
        for run in list(parrafo.runs):
            run.text = ""
        if parrafo.runs:
            parrafo.runs[0].text = texto
        else:
            parrafo.add_run(texto)


def _reemplazar_en_documento(doc: Document, reemplazos: dict[str, str]) -> None:
    for p in doc.paragraphs:
        _reemplazar_texto_parrafo(p, reemplazos)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    _reemplazar_texto_parrafo(p, reemplazos)


def _normalizar_etiqueta(texto: str) -> str:
    return texto.replace("​", "").replace("\xa0", " ").strip().upper()


def _set_por_etiqueta(doc: Document, etiqueta_buscada: str, valor: str, exacto: bool = True) -> bool:
    """Busca en TODAS las tablas una fila cuya primera celda coincida con
    `etiqueta_buscada` (comparación exacta por defecto) y reemplaza el
    texto de la segunda celda. Es necesario para columnas como
    'DIRECCION:' y 'DIRECCION TERRITORIAL:' que en la plantilla comparten
    el mismo marcador de texto pero deben llevar valores distintos --
    aquí se distinguen por la etiqueta de la fila, no por el marcador."""
    objetivo = _normalizar_etiqueta(etiqueta_buscada)
    for t in doc.tables:
        for row in t.rows:
            if len(row.cells) < 2:
                continue
            etiqueta = _normalizar_etiqueta(row.cells[0].text)
            coincide = (etiqueta == objetivo) if exacto else (objetivo in etiqueta)
            if coincide:
                celda_valor = row.cells[1]
                p = celda_valor.paragraphs[0]
                for run in list(p.runs):
                    run.text = ""
                if p.runs:
                    p.runs[0].text = valor or "—"
                else:
                    p.add_run(valor or "—")
                return True
    return False


def _insertar_imagen_en_marcador(doc: Document, marcador: str, contenido_imagen: bytes | None, ancho_mm: float = 140) -> bool:
    if not contenido_imagen:
        return False
    for p in doc.paragraphs:
        if marcador in p.text:
            for run in list(p.runs):
                run.text = ""
            run = p.runs[0] if p.runs else p.add_run()
            try:
                run.add_picture(io.BytesIO(contenido_imagen), width=Mm(ancho_mm))
            except Exception:
                return False
            return True
    return False


def _quitar_bordes_tabla(tabla) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tbl = tabla._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for borde in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{borde}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)


def _insertar_fotos_grid(doc: Document, marcador: str, fotos: list[bytes], columnas: int = 6,
                          ancho_fila_mm: float = 150, ancho_maximo_mm: float = 70) -> bool:
    fotos = [f for f in fotos if f]
    if not fotos:
        return False
    parrafo_marcador = None
    for p in doc.paragraphs:
        if marcador in p.text:
            parrafo_marcador = p
            break
    if parrafo_marcador is None:
        return False
    for run in list(parrafo_marcador.runs):
        run.text = ""

    columnas_usadas = min(columnas, len(fotos))
    ancho_mm = min(ancho_maximo_mm, ancho_fila_mm / columnas_usadas)
    filas = [fotos[i:i + columnas] for i in range(0, len(fotos), columnas)]

    tabla = doc.add_table(rows=len(filas), cols=columnas)
    tabla.autofit = True
    _quitar_bordes_tabla(tabla)

    for fila_idx, fila_fotos in enumerate(filas):
        for col_idx in range(columnas):
            celda = tabla.cell(fila_idx, col_idx)
            parrafo_celda = celda.paragraphs[0]
            if col_idx < len(fila_fotos):
                run = parrafo_celda.add_run()
                try:
                    run.add_picture(io.BytesIO(fila_fotos[col_idx]), width=Mm(ancho_mm))
                except Exception:
                    parrafo_celda.add_run("(foto no disponible)")

    parrafo_marcador._p.addnext(tabla._tbl)
    return True


def generar_acta(datos: dict, fotos_inmueble: list[bytes], foto_sagrilaft: bytes | None,
                  foto_estimado_renta: bytes | None = None) -> bytes:
    """`datos` trae, ya con las correcciones manuales aplicadas:
    fmi, territorial, tipo_bien, tipo_contrato, direccion,
    arrendatario_nombre, id_siglas, id_numero, id_ciudad, descripcion."""
    doc = Document(RUTA_PLANTILLA)

    # Marcadores que aparecen una sola vez en la plantilla: reemplazo
    # global normal.
    reemplazos = {
        "##fmi_estimado_renta##": datos.get("fmi", "—"),
        "##tipo_bien_estimado_renta##": datos.get("tipo_bien", "—"),
        "##descripcion_estimado_renta##": datos.get("descripcion") or "—",
        "##carpeta_arrendatario_nombre##": datos.get("arrendatario_nombre") or "—",
        "##contrato##": datos.get("tipo_contrato", "—"),
        "##numero_consecutivo##": "",  # lo asigna la Coordinación/Comité
        "##cedula/nit_siglas##": datos.get("id_siglas") or "C.C.",
        "##cedula/nit_numero##": datos.get("id_numero") or "—",
        "##cedula/nit_ciudad##": datos.get("id_ciudad") or "—",
    }
    _reemplazar_en_documento(doc, reemplazos)

    # "DIRECCION:" y "DIRECCION TERRITORIAL:" comparten el mismo texto de
    # marcador en la plantilla (##direccion_estimado_renta##) pero deben
    # llevar valores distintos -- se resuelven por la etiqueta de la fila,
    # no por el marcador, para no pisar uno con el otro.
    _set_por_etiqueta(doc, "DIRECCION TERRITORIAL:", datos.get("territorial") or "—")
    _set_por_etiqueta(doc, "DIRECCION:", datos.get("direccion") or "—")

    _insertar_fotos_grid(doc, "##foto_estimado_renta##", fotos_inmueble, columnas=6)
    _insertar_imagen_en_marcador(doc, "##carpeta_foto_aprobado_sagrilaft##", foto_sagrilaft)
    _insertar_imagen_en_marcador(doc, "##archivo_foto_preaprobado_poliza##", foto_estimado_renta)

    # Reemplaza cualquier marcador de imagen que haya quedado sin llenar
    # con el texto "No hay imagen para este documento", en vez de dejar
    # el marcador crudo visible en el documento final.
    for p in doc.paragraphs:
        if "##" in p.text:
            _reemplazar_texto_parrafo(p, {
                "##carpeta_foto_aprobado_sagrilaft##": SIN_IMAGEN,
                "##archivo_foto_preaprobado_poliza##": SIN_IMAGEN,
                "##foto_estimado_renta##": SIN_IMAGEN,
            })

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
