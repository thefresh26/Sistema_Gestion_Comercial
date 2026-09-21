"""
Tipo de documento: Acta de Alcance.

Complementa una Acta de Comité de Adjudicación ANTIGUA (subida como
documento fuente, en .docx o .pdf) con lo que el "Modelo Vigente" exige
agregar: código ActiBid, oferente ganador y su identificación, y el
monto de la puja ganadora -- estos últimos salen de la base de datos de
negocio (subastas), igual que el script original.

Valor catastral, PMV y "¿superó el precio base?" no se pueden sacar
automáticamente todavía (esos vivían en un Excel de paquete que este
sistema web no lee) y quedan marcados como "[POR COMPLETAR]", igual que
en el script original -- están listados en los "pendientes" de la Acta.
"""
from __future__ import annotations

import copy
import io
import os
import re
from datetime import date

import docx
from docx.shared import Mm

from core import db

RAIZ_PROYECTO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RUTA_PLANTILLA = os.path.join(RAIZ_PROYECTO, "word_templates", "ACTA_DE_ALCANCE.docx")
CARPETA_FIRMAS = os.path.join(RAIZ_PROYECTO, "word_templates", "firmas")
FIRMAS = {
    "heidy": os.path.join(CARPETA_FIRMAS, "firma_heidy.png"),
    "jairo": os.path.join(CARPETA_FIRMAS, "firma_jairo.png"),
    "jose": os.path.join(CARPETA_FIRMAS, "firma_jose.png"),
    "julio": os.path.join(CARPETA_FIRMAS, "firma_julio.png"),
}
VALOR_POR_COMPLETAR = "[POR COMPLETAR]"

TIPOS_DOCUMENTO_FUENTE = {
    "acta_vieja": "Acta de Comité de Adjudicación original (.docx o .pdf)",
}

CAMPOS_EDITABLES = [
    ("codigo_inmueble", "Código ActiBid"),
    ("nombre_ganador", "Oferente ganador"),
    ("id_ganador", "Cédula/NIT del ganador"),
    ("monto_ganador", "Monto adjudicado"),
]


# ---------------------------------------------------------------------
# Lectura del Acta vieja (misma lógica que el script original)
# ---------------------------------------------------------------------

def _normalizar_fecha_texto(texto: str) -> str:
    m = re.match(r"(\d{1,2})\s+de\s+(\w+)\s+(?:de\s+)?(\d{4})", texto.strip(), re.IGNORECASE)
    if m:
        dia, mes, anio = m.groups()
        return f"{dia} de {mes} de {anio}"
    return texto.strip()


def _fecha_hoy_es() -> str:
    meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
             "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
    hoy = date.today()
    return f"{hoy.day:02d} de {meses[hoy.month - 1]} de {hoy.year}"


def _listar_con_y(items: list[str]) -> str:
    items = [i for i in items if i and i != "—"]
    if not items:
        return "—"
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " y " + items[-1]


def _extraer_campos_comunes(texto_parrafos: str, tablas: list) -> dict:
    m_num = re.search(r"ADJUDICAC\w*.*?No\.?\s*0*(\d+)", texto_parrafos, re.IGNORECASE | re.DOTALL)
    if not m_num:
        m_num = re.search(r"\bNo\.?\s*0*(\d{2,4})\b", texto_parrafos)
    numero_acta_vieja = m_num.group(1) if m_num else "—"

    fecha_acta_vieja = "—"
    for tabla in tablas:
        for fila in tabla:
            celdas = [(c or "").strip() for c in fila]
            if celdas and celdas[0].strip().lower() == "fecha" and len(celdas) > 1:
                candidato = celdas[1].strip()
                if candidato and candidato.lower() != "fecha":
                    fecha_acta_vieja = candidato
                    break
        if fecha_acta_vieja != "—":
            break

    m_paq = re.search(r"Paquete\s*No\.?\s*(\d+)", texto_parrafos, re.IGNORECASE)
    numero_paquete = m_paq.group(1) if m_paq else "—"

    inmuebles = []
    for tabla in tablas:
        if not tabla:
            continue
        idx_encabezado = None
        fila_encabezado = None
        for i, fila in enumerate(tabla[:3]):
            celdas_enc = [(c or "").strip().lower() for c in fila]
            if "fmi" in celdas_enc and any("tipolog" in c or c == "id" for c in celdas_enc):
                idx_encabezado = i
                fila_encabezado = fila
                break
        if idx_encabezado is None:
            continue

        filas_encabezado = [fila_encabezado]
        j = idx_encabezado + 1
        while j < len(tabla) and (j - idx_encabezado) <= 5:
            fila_j = tabla[j]
            celdas_j = [(c or "").strip() for c in fila_j]
            no_vacias = [c for c in celdas_j if c]
            if not no_vacias:
                filas_encabezado.append(fila_j)
                j += 1
                continue
            tiene_digitos = any(ch.isdigit() for ch in " ".join(no_vacias))
            if not tiene_digitos and all(len(c) <= 20 for c in no_vacias):
                filas_encabezado.append(fila_j)
                j += 1
                continue
            break
        idx_datos_inicio = j

        num_columnas = max(len(f) for f in filas_encabezado)
        encabezado = []
        for col_idx in range(num_columnas):
            partes = []
            for fila_h in filas_encabezado:
                if col_idx < len(fila_h) and fila_h[col_idx]:
                    partes.append(str(fila_h[col_idx]).strip())
            encabezado.append(" ".join(partes).strip().lower())

        idx = {nombre: i for i, nombre in enumerate(encabezado) if nombre}

        def _col(claves):
            for clave in claves:
                for nombre_col, i in idx.items():
                    if clave in nombre_col:
                        return i
            return None

        col_id = _col(["id"])
        col_fmi = _col(["fmi"])
        col_tipo = _col(["tipolog"])
        col_dir = _col(["direcci"])
        col_mun = _col(["municipio"])
        col_ocup = _col(["ocupaci"])
        col_jur = _col(["jur"])
        col_fsub = _col(["subasta"])

        for fila in tabla[idx_datos_inicio:]:
            celdas = [(c or "").replace("\n", " ").strip() for c in fila]
            if not any(celdas):
                continue
            valor_id = celdas[col_id] if col_id is not None and col_id < len(celdas) else ""
            valor_fmi = celdas[col_fmi] if col_fmi is not None and col_fmi < len(celdas) else ""
            if not valor_id and not valor_fmi:
                continue

            def _val(col, quitar_espacios=False):
                if col is None or col >= len(celdas) or not celdas[col]:
                    return "—"
                valor = celdas[col]
                return valor.replace(" ", "") if quitar_espacios else valor

            inmuebles.append({
                "id": _val(col_id), "fmi": _val(col_fmi, quitar_espacios=True),
                "tipologia": _val(col_tipo), "direccion": _val(col_dir),
                "municipio": _val(col_mun), "estado_ocupacion": _val(col_ocup),
                "estado_juridico": _val(col_jur), "fecha_subasta": _val(col_fsub),
            })
        if inmuebles:
            break

    return {
        "numero_acta_vieja": numero_acta_vieja,
        "fecha_acta_vieja": fecha_acta_vieja,
        "numero_paquete": numero_paquete,
        "inmuebles": inmuebles,
    }


def _leer_acta_vieja(contenido: bytes, nombre_archivo: str) -> dict:
    if nombre_archivo.lower().endswith(".pdf"):
        import pdfplumber
        texto_partes, tablas = [], []
        with pdfplumber.open(io.BytesIO(contenido)) as pdf:
            for pagina in pdf.pages:
                texto_partes.append(pagina.extract_text() or "")
                tablas.extend(pagina.extract_tables())
        return _extraer_campos_comunes("\n".join(texto_partes), tablas)

    doc = docx.Document(io.BytesIO(contenido))
    texto_parrafos = "\n".join(p.text for p in doc.paragraphs)
    tablas = [[[c.text for c in fila.cells] for fila in tabla.rows] for tabla in doc.tables]
    return _extraer_campos_comunes(texto_parrafos, tablas)


# ---------------------------------------------------------------------
# Datos de la base de datos de negocio (subastas)
# ---------------------------------------------------------------------

def _resolver_auction_uuid(conn, identificador: str) -> str | None:
    if re.match(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", identificador.lower()):
        return identificador
    fila = conn.execute(
        "SELECT id FROM polybid.auctions WHERE code = %(c)s LIMIT 1", {"c": identificador},
    ).fetchone()
    if fila:
        return str(fila["id"])

    filas = conn.execute(
        """
        SELECT id, grupo_id FROM mst_inmuebles
        WHERE UPPER(numero_matricula) = UPPER(%(id)s) OR UPPER(codigo) = UPPER(%(id)s)
           OR UPPER(codigo_grupo) = UPPER(%(id)s) OR UPPER(referencia) = UPPER(%(id)s)
        """,
        {"id": identificador},
    ).fetchall()
    if not filas:
        return None
    inmueble_ids = sorted({f["id"] for f in filas if f["id"] is not None})
    grupo_ids = sorted({f["grupo_id"] for f in filas if f["grupo_id"] is not None})

    candidatos = conn.execute(
        """
        SELECT DISTINCT COALESCE(a.id, psv.auction_id) AS auction_id,
               COALESCE(a.start_date, psv.fecha_inicio) AS start_date
        FROM polibid_subastas_v2 psv
        LEFT JOIN polybid.auctions a ON a.id = psv.auction_id
        WHERE psv.inmueble_id = ANY(%(inm)s)
           OR (psv.grupo_id IS NOT NULL AND psv.grupo_id = ANY(%(grp)s))
        ORDER BY start_date DESC NULLS LAST
        """,
        {"inm": inmueble_ids or [-1], "grp": grupo_ids or [-1]},
    ).fetchall()
    if not candidatos:
        return None
    return str(candidatos[0]["auction_id"])


def _ganador_de_subasta(conn, auction_uuid: str) -> dict | None:
    """Misma lógica que `fetch_informe` en core/core.py: primero la puja
    marcada como 'WINNING'; si aún no existe (subastas activas), la mejor
    oferta vigente por monto."""
    fila = conn.execute(
        """
        SELECT b.amount, ct.nombre_principal, ct.identificacion_numero
        FROM polybid.auction_bids b
        INNER JOIN polybid.auction_participants p
            ON p.auction_id = b.auction_id AND p.client_id = b.client_id
        LEFT JOIN polibid_credentials pc ON pc.client_id = b.client_id
        LEFT JOIN contact_terceros ct ON ct.id = pc.contact_tercero_id
        WHERE b.auction_id = %(id)s::uuid AND b.status = 'WINNING'
        LIMIT 1
        """,
        {"id": auction_uuid},
    ).fetchone()
    if not fila:
        fila = conn.execute(
            """
            SELECT b.amount, ct.nombre_principal, ct.identificacion_numero
            FROM polybid.auction_bids b
            INNER JOIN polybid.auction_participants p
                ON p.auction_id = b.auction_id AND p.client_id = b.client_id
            LEFT JOIN polibid_credentials pc ON pc.client_id = b.client_id
            LEFT JOIN contact_terceros ct ON ct.id = pc.contact_tercero_id
            WHERE b.auction_id = %(id)s::uuid
            ORDER BY b.amount DESC, b.created_at DESC
            LIMIT 1
            """,
            {"id": auction_uuid},
        ).fetchone()
    return fila


def _codigo_inmueble_de_subasta(conn, auction_uuid: str) -> str:
    link = conn.execute(
        """SELECT inmueble_id, grupo_id FROM polibid_subastas_v2
           WHERE auction_id = %(id)s::uuid ORDER BY id DESC LIMIT 1""",
        {"id": auction_uuid},
    ).fetchone()
    if not link:
        return "—"
    if link["inmueble_id"]:
        fila = conn.execute(
            "SELECT codigo, codigo_grupo FROM mst_inmuebles WHERE id = %(id)s",
            {"id": link["inmueble_id"]},
        ).fetchone()
    elif link["grupo_id"]:
        fila = conn.execute(
            """SELECT codigo, codigo_grupo FROM mst_inmuebles WHERE grupo_id = %(gid)s
               ORDER BY es_padre DESC, id LIMIT 1""",
            {"gid": link["grupo_id"]},
        ).fetchone()
    else:
        fila = None
    if not fila:
        return "—"
    return fila["codigo_grupo"] or fila["codigo"] or "—"


def _fmt_numero(valor) -> str:
    if not valor:
        return "0"
    try:
        return f"{int(valor):,}".replace(",", ".")
    except Exception:
        return str(valor)


def obtener_datos_bd(identificador: str) -> dict:
    with db.get_conn_negocio() as conn:
        auction_uuid = _resolver_auction_uuid(conn, identificador)
        if not auction_uuid:
            return {"codigo_inmueble": "—", "nombre_ganador": "—", "id_ganador": "—", "monto_ganador": "0"}
        ganador = _ganador_de_subasta(conn, auction_uuid)
        codigo_inmueble = _codigo_inmueble_de_subasta(conn, auction_uuid)
    return {
        "codigo_inmueble": codigo_inmueble,
        "nombre_ganador": (ganador["nombre_principal"] or "—").upper() if ganador else "—",
        "id_ganador": str(ganador["identificacion_numero"] or "—") if ganador else "—",
        "monto_ganador": _fmt_numero(ganador["amount"]) if ganador else "0",
    }


# ---------------------------------------------------------------------
# Generación del .docx
# ---------------------------------------------------------------------

def _duplicar_fila(tabla, indice: int, veces: int):
    tr_referencia = tabla.rows[indice]._tr
    ultimo_tr = tr_referencia
    for _ in range(veces):
        nuevo_tr = copy.deepcopy(tr_referencia)
        ultimo_tr.addnext(nuevo_tr)
        ultimo_tr = nuevo_tr
    return list(tabla.rows[indice: indice + 1 + veces])


def _reemplazar_texto_parrafo(parrafo, mapa: dict) -> None:
    texto_completo = parrafo.text
    if not texto_completo or not any(m in texto_completo for m in mapa):
        return
    nuevo_texto = texto_completo
    for marcador, valor in mapa.items():
        nuevo_texto = nuevo_texto.replace(marcador, str(valor))
    if parrafo.runs:
        parrafo.runs[0].text = nuevo_texto
        for run in parrafo.runs[1:]:
            run.text = ""
    else:
        parrafo.add_run(nuevo_texto)


def _insertar_firma(celda, ruta_imagen: str, ancho_mm: float = 32) -> bool:
    if not os.path.isfile(ruta_imagen):
        return False
    for p in celda.paragraphs:
        if "##firma_" in p.text:
            for run in p.runs:
                run.text = ""
            run = p.runs[0] if p.runs else p.add_run()
            try:
                run.add_picture(ruta_imagen, width=Mm(ancho_mm))
            except Exception:
                return False
            return True
    return False


def _generar_docx(acta_vieja: dict, bd: dict, num_consecutivo: int) -> tuple[bytes, list[str]]:
    doc = docx.Document(RUTA_PLANTILLA)
    pendientes: list[str] = []

    inmuebles = acta_vieja["inmuebles"] or [{
        "id": "—", "fmi": "—", "tipologia": "—", "direccion": "—", "municipio": "—",
        "estado_ocupacion": "—", "estado_juridico": "—", "fecha_subasta": "—",
    }]
    fecha_normalizada = _normalizar_fecha_texto(acta_vieja["fecha_acta_vieja"])

    if len(inmuebles) == 1:
        descripcion = ", ".join(
            x for x in [inmuebles[0]["tipologia"], inmuebles[0]["direccion"], inmuebles[0]["municipio"]]
            if x and x != "—"
        )
    else:
        descripcion = "; ".join(
            f"{i['tipologia']} {i['direccion']}".strip()
            for i in inmuebles if i.get("direccion") and i["direccion"] != "—"
        )
    paquete_y_direccion = f"Paquete No. {acta_vieja['numero_paquete']}"
    if descripcion:
        paquete_y_direccion += f" ({descripcion})"

    valor_catastral_txt = VALOR_POR_COMPLETAR
    precio_minimo_txt = VALOR_POR_COMPLETAR
    supero_txt = VALOR_POR_COMPLETAR

    numero_paquete = acta_vieja.get("numero_paquete", "—")
    cronograma_txt = f"Cronograma Paquete No. {numero_paquete}" if numero_paquete != "—" else VALOR_POR_COMPLETAR

    mapa_texto = {
        "##num_consecutivo_actual##": str(num_consecutivo),
        "##num_consecutivo_acta_existente##": f"No. {acta_vieja['numero_acta_vieja']}",
        "##fecha_acta_vieja##": f"del {fecha_normalizada}" if fecha_normalizada != "—" else "—",
        "##fmis_acta_aterior##": _listar_con_y([i["fmi"] for i in inmuebles]),
        "##paquete_y_direccion##": paquete_y_direccion,
        "##fecha_hoy##": _fecha_hoy_es(),
        "##codigo_inmueble##": bd["codigo_inmueble"],
        "##fmi##": _listar_con_y([i["fmi"] for i in inmuebles]),
        "##fmi_catastral##": valor_catastral_txt,
        "##fmi_catastral ##": valor_catastral_txt,
        "##cronograma_inmueble##": cronograma_txt,
        "##tipologia##": inmuebles[0]["tipologia"],
        "##direccion_fmi##": inmuebles[0]["direccion"],
        "##municipio_fmi##": inmuebles[0]["municipio"],
        "##estado##": inmuebles[0]["estado_ocupacion"],
        "##estado_juridico##": inmuebles[0]["estado_juridico"],
        "##fecha_subasta##": inmuebles[0]["fecha_subasta"],
        "##precio_minimo##": precio_minimo_txt,
        "##oferta_ganadora##": f"$ {bd['monto_ganador']}",
        "##¿supero_precio_base?##": supero_txt,
        "##nombre_participante##": bd["nombre_ganador"],
        "##nit/cedula##": bd["id_ganador"],
        "##puja_ganadora##": f"$ {bd['monto_ganador']}",
        "##valor_cerrado_subasta##": f"$ {bd['monto_ganador']}",
    }

    if bd["codigo_inmueble"] == "—":
        pendientes.append("Código ActiBid: no se encontró en la base de datos, revisar a mano.")
    if bd["nombre_ganador"] == "—":
        pendientes.append("Oferente ganador: no se encontró en la base de datos, revisar a mano.")
    pendientes.append("Valor catastral vigente: este sistema todavía no lee el Excel del paquete, completar a mano.")
    pendientes.append("Precio Mínimo de Venta (PMV): este sistema todavía no lee el Excel del paquete, completar a mano.")
    pendientes.append("¿Superó el precio base?: depende del PMV, completar a mano.")
    if cronograma_txt == VALOR_POR_COMPLETAR:
        pendientes.append("Cronograma del inmueble: no se encontró el número de paquete, revisar a mano.")
    pendientes.append("Fecha límite del compromiso de la Secretaría Técnica: revisar si la de la plantilla aplica.")

    def _ajustes_texto_literal(texto: str) -> str:
        texto = texto.replace(
            "Acta No. ##num_consecutivo_acta_existente##", "Acta ##num_consecutivo_acta_existente##",
        )
        texto = texto.replace("acta inicial No. 019", f"acta inicial No. {acta_vieja['numero_acta_vieja']}")
        texto = texto.replace("Acta inicial No. 019", f"Acta inicial No. {acta_vieja['numero_acta_vieja']}")
        return texto

    def _aplicar_ajustes_parrafo(p) -> None:
        texto_original = p.text
        texto_nuevo = _ajustes_texto_literal(texto_original)
        if texto_nuevo != texto_original and p.runs:
            p.runs[0].text = texto_nuevo
            for r in p.runs[1:]:
                r.text = ""

    for p in doc.paragraphs:
        _aplicar_ajustes_parrafo(p)
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    _aplicar_ajustes_parrafo(p)

    tabla_inmuebles = None
    for tabla in doc.tables:
        encabezado = [c.text.strip().lower() for c in tabla.rows[0].cells] if tabla.rows else []
        if "fmi" in encabezado and "tipología" in encabezado:
            tabla_inmuebles = tabla
            break

    if tabla_inmuebles is not None and len(inmuebles) > 1:
        filas = _duplicar_fila(tabla_inmuebles, 1, len(inmuebles) - 1)
        for fila, inm in zip(filas, inmuebles):
            fila.cells[0].text = inm["id"]
            fila.cells[1].text = inm["fmi"]
            fila.cells[2].text = inm["tipologia"]
            fila.cells[3].text = inm["direccion"]
            fila.cells[4].text = inm["municipio"]
            fila.cells[5].text = inm["estado_ocupacion"]
            fila.cells[6].text = inm["estado_juridico"]
            fila.cells[7].text = inm["fecha_subasta"]
    elif tabla_inmuebles is not None:
        fila = tabla_inmuebles.rows[1]
        fila.cells[0].text = inmuebles[0]["id"]

    for p in doc.paragraphs:
        _reemplazar_texto_parrafo(p, mapa_texto)
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if "##firma_" in celda.text:
                    continue
                for p in celda.paragraphs:
                    _reemplazar_texto_parrafo(p, mapa_texto)

    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for persona, ruta_img in FIRMAS.items():
                    marcador = f"##firma_{persona}##"
                    if marcador in celda.text:
                        if not _insertar_firma(celda, ruta_img):
                            pendientes.append(f"Firma de {persona}: no se pudo insertar la imagen, revisar a mano.")

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue(), pendientes


def generar(fmi: str) -> tuple[bytes, str, list[str]]:
    """`fmi` aquí es el identificador para buscar en la base de datos de
    negocio (FMI, código de subasta o de inmueble/unidad). El Acta vieja
    debe estar subida como documento fuente 'acta_vieja' para este caso."""
    doc_acta_vieja = db.obtener_ultimo_documento(fmi, "acta_vieja")
    if not doc_acta_vieja:
        raise ValueError(
            "Falta subir el Acta de Comité de Adjudicación original (.docx o .pdf) "
            "como documento fuente antes de poder generar el Acta de Alcance."
        )
    acta_vieja = _leer_acta_vieja(doc_acta_vieja["contenido"], doc_acta_vieja["nombre_archivo"])

    identificador_bd = fmi
    if acta_vieja["inmuebles"] and acta_vieja["inmuebles"][0]["fmi"] != "—":
        identificador_bd = acta_vieja["inmuebles"][0]["fmi"]

    bd = obtener_datos_bd(identificador_bd)
    consecutivo = db.siguiente_consecutivo("acta_alcance", inicial=81)

    contenido, pendientes = _generar_docx(acta_vieja, bd, consecutivo)
    codigo_archivo = re.sub(r"[^A-Za-z0-9]+", "_", fmi)
    nombre_archivo = f"ACTA_DE_ALCANCE_{consecutivo}_{codigo_archivo}.docx"

    db.upsert_caso(fmi, {**bd, "estado": "con_observaciones" if pendientes else "completo", "pendientes": pendientes})
    return contenido, nombre_archivo, pendientes
